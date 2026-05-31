#!/usr/bin/env python3
"""
Convert Markdown Report to Word Document
Generates lifeosfinalreport.docx from lifeosfinalreport.md
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

def set_cell_border(cell, **kwargs):
    """Set cell borders"""
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        if edge in kwargs:
            edge_el = OxmlElement(f'w:{edge}')
            edge_el.set(qn('w:val'), 'single')
            edge_el.set(qn('w:sz'), '12')
            edge_el.set(qn('w:space'), '0')
            edge_el.set(qn('w:color'), 'CCCCCC')
            tcBorders.append(edge_el)
    tcPr.append(tcBorders)

def read_markdown_file(filepath):
    """Read markdown file"""
    with open(filepath, 'r', encoding='utf-8') as f:
        return f.read()

def create_word_document(markdown_content, output_path):
    """Convert markdown to Word document"""
    doc = Document()

    # Set document margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Split by lines
    lines = markdown_content.split('\n')

    current_list = None
    in_code_block = False
    code_language = None
    code_lines = []

    i = 0
    while i < len(lines):
        line = lines[i]

        # Handle code blocks
        if line.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_language = line[3:].strip()
                code_lines = []
            else:
                in_code_block = False
                if code_lines:
                    code_para = doc.add_paragraph()
                    code_para.paragraph_format.left_indent = Inches(0.5)
                    code_para.paragraph_format.space_before = Pt(6)
                    code_para.paragraph_format.space_after = Pt(6)

                    code_text = '\n'.join(code_lines)
                    code_run = code_para.add_run(code_text)
                    code_run.font.name = 'Courier New'
                    code_run.font.size = Pt(9)
                    code_run.font.color.rgb = RGBColor(100, 100, 100)
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Handle headings
        if line.startswith('# '):
            heading_text = line[2:].strip()
            heading = doc.add_heading(heading_text, level=1)
            heading.paragraph_format.space_before = Pt(12)
            heading.paragraph_format.space_after = Pt(12)
            i += 1
            continue

        if line.startswith('## '):
            heading_text = line[3:].strip()
            heading = doc.add_heading(heading_text, level=2)
            heading.paragraph_format.space_before = Pt(10)
            heading.paragraph_format.space_after = Pt(10)
            i += 1
            continue

        if line.startswith('### '):
            heading_text = line[4:].strip()
            heading = doc.add_heading(heading_text, level=3)
            heading.paragraph_format.space_before = Pt(8)
            heading.paragraph_format.space_after = Pt(8)
            i += 1
            continue

        if line.startswith('#### '):
            heading_text = line[5:].strip()
            heading = doc.add_heading(heading_text, level=4)
            heading.paragraph_format.space_before = Pt(6)
            heading.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        # Handle tables
        if line.startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].startswith('|'):
                table_lines.append(lines[i])
                i += 1

            if len(table_lines) > 2:
                # Parse table
                rows = []
                for line in table_lines:
                    if '|' in line:
                        cells = [cell.strip() for cell in line.split('|')[1:-1]]
                        rows.append(cells)

                if rows:
                    # Create table
                    num_cols = len(rows[0])
                    table = doc.add_table(rows=len(rows), cols=num_cols)
                    table.style = 'Light Grid Accent 1'

                    # Fill table
                    for row_idx, row_data in enumerate(rows):
                        for col_idx, cell_data in enumerate(row_data):
                            cell = table.rows[row_idx].cells[col_idx]
                            # Skip separator row
                            if '---' not in cell_data:
                                cell.text = cell_data
                                # Bold header row
                                if row_idx == 0:
                                    for paragraph in cell.paragraphs:
                                        for run in paragraph.runs:
                                            run.font.bold = True
            continue

        # Handle lists
        if line.startswith('- '):
            if current_list != 'unordered':
                current_list = 'unordered'

            list_text = line[2:].strip()
            p = doc.add_paragraph(list_text, style='List Bullet')
            p.paragraph_format.left_indent = Inches(0.5)
            i += 1
            continue

        # Handle ordered lists
        if re.match(r'^\d+\. ', line):
            if current_list != 'ordered':
                current_list = 'ordered'

            list_text = re.sub(r'^\d+\. ', '', line).strip()
            p = doc.add_paragraph(list_text, style='List Number')
            p.paragraph_format.left_indent = Inches(0.5)
            i += 1
            continue

        current_list = None

        # Handle horizontal rules
        if line.startswith('---'):
            doc.add_paragraph()
            i += 1
            continue

        # Handle empty lines
        if not line.strip():
            doc.add_paragraph()
            i += 1
            continue

        # Handle regular paragraphs
        if line.strip():
            # Parse inline formatting
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*|__.*?__|`.*?`|\*.*?\*|_.*?_|\[.*?\]\(.*?\))', line)

            for part in parts:
                if not part:
                    continue

                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                elif part.startswith('__') and part.endswith('__'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                elif part.startswith('`') and part.endswith('`'):
                    run = p.add_run(part[1:-1])
                    run.font.name = 'Courier New'
                    run.font.size = Pt(10)
                elif part.startswith('*') and part.endswith('*') and len(part) > 1:
                    run = p.add_run(part[1:-1])
                    run.font.italic = True
                elif part.startswith('_') and part.endswith('_') and len(part) > 1:
                    run = p.add_run(part[1:-1])
                    run.font.italic = True
                elif part.startswith('[') and '](' in part:
                    match = re.match(r'\[(.*?)\]\((.*?)\)', part)
                    if match:
                        text = match.group(1)
                        url = match.group(2)
                        run = p.add_run(text)
                        run.font.underline = True
                        run.font.color.rgb = RGBColor(0, 0, 255)
                else:
                    p.add_run(part)

            p.paragraph_format.space_after = Pt(6)

        i += 1

    # Save document
    doc.save(output_path)
    print(f"✓ Word document created: {output_path}")

if __name__ == '__main__':
    markdown_path = 'Documentation/lifeosfinalreport.md'
    docx_path = 'Documentation/lifeosfinalreport.docx'

    print("Converting Markdown to Word Document...")
    print(f"Reading from: {markdown_path}")

    try:
        content = read_markdown_file(markdown_path)
        print(f"Markdown file loaded: {len(content)} characters")

        create_word_document(content, docx_path)
        print(f"✓ Conversion complete!")
        print(f"Output file: {docx_path}")

    except Exception as e:
        print(f"✗ Error: {str(e)}")
        import traceback
        traceback.print_exc()

